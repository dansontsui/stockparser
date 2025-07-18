#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
每日除息資料檢查器
自動檢查庫存股票的除息資料並計算配息收益
"""

import requests
import json
import pandas as pd
from datetime import datetime, timedelta
import urllib3
import os
from stock_inventory_system import StockInventorySystem
from typing import Dict, List, Optional, Tuple

# 禁用SSL警告
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)


class DailyDividendChecker:
    """每日除息資料檢查器"""
    
    def __init__(self, inventory_file: str = "stock_inventory.xlsx"):
        """
        初始化除息檢查器

        Args:
            inventory_file: 庫存Excel檔案路徑
        """
        self.inventory_file = inventory_file
        self.stock_system = StockInventorySystem(inventory_file)
        # 櫃買中心API (興櫃、上櫃股票)
        self.tpex_url = "https://www.tpex.org.tw/web/stock/exright/dailyquo/exDailyQ_result.php"
        # 證交所API (上市股票) - 除權息資料
        self.tse_url = "https://www.twse.com.tw/rwd/zh/exRight/TWT49U"
        
    def get_date_range_roc(self) -> Tuple[str, str]:
        """
        取得前一個月的日期範圍 (民國年格式)
        
        Returns:
            (start_date, end_date) 民國年格式 例如: ("113/12/01", "113/12/31")
        """
        today = datetime.now()
        
        # 計算前一個月
        if today.month == 1:
            # 如果是1月，前一個月是去年12月
            prev_month = 12
            prev_year = today.year - 1
        else:
            prev_month = today.month - 1
            prev_year = today.year
        
        # 前一個月的第一天
        start_date = datetime(prev_year, prev_month, 1)
        
        # 前一個月的最後一天
        if prev_month == 12:
            next_month_first = datetime(prev_year + 1, 1, 1)
        else:
            next_month_first = datetime(prev_year, prev_month + 1, 1)
        end_date = next_month_first - timedelta(days=1)
        
        # 轉換為民國年格式
        start_roc = f"{start_date.year - 1911}/{start_date.month:02d}/{start_date.day:02d}"
        end_roc = f"{end_date.year - 1911}/{end_date.month:02d}/{end_date.day:02d}"
        
        return start_roc, end_roc
    
    def fetch_dividend_data_from_tpex(self, start_date: str, end_date: str) -> Optional[pd.DataFrame]:
        """
        從櫃買中心獲取除息資料

        Args:
            start_date: 開始日期 (民國年格式)
            end_date: 結束日期 (民國年格式)

        Returns:
            除息資料DataFrame，失敗時返回None
        """
        url = f"{self.tpex_url}?l=zh-tw&d={start_date}&ed={end_date}"

        print(f"🔍 從櫃買中心獲取除息資料...")
        print(f"📅 查詢期間: {start_date} ~ {end_date}")
        print(f"🌐 URL: {url}")

        try:
            # 發送請求
            response = requests.get(url, timeout=30, verify=False)

            if response.status_code != 200:
                print(f"❌ 櫃買中心HTTP請求失敗: {response.status_code}")
                return None

            # 解析JSON
            json_data = response.json()

            # 遞迴搜尋 fields 和 data
            def search_fields_data(obj, path=""):
                if isinstance(obj, dict):
                    if 'fields' in obj and 'data' in obj:
                        return obj['fields'], obj['data']

                    for key, value in obj.items():
                        new_path = f"{path}.{key}" if path else key
                        result = search_fields_data(value, new_path)
                        if result[0] is not None:
                            return result

                elif isinstance(obj, list):
                    for i, item in enumerate(obj):
                        new_path = f"{path}[{i}]" if path else f"[{i}]"
                        result = search_fields_data(item, new_path)
                        if result[0] is not None:
                            return result

                return None, None

            fields, data = search_fields_data(json_data)

            if fields is None or data is None:
                print("❌ 櫃買中心未找到除息資料結構")
                return None

            # 創建DataFrame
            df = pd.DataFrame(data, columns=fields)
            df['資料來源'] = '櫃買中心'
            print(f"✅ 櫃買中心成功獲取 {len(df)} 筆除息資料")

            # 儲存櫃買中心資料為Excel
            self._save_raw_data_to_excel(df, '櫃買中心', start_date, end_date)

            return df

        except Exception as e:
            print(f"❌ 櫃買中心獲取除息資料失敗: {e}")
            return None

    def fetch_dividend_data_from_tse(self, start_date: str, end_date: str) -> Optional[pd.DataFrame]:
        """
        從證交所獲取除息資料

        Args:
            start_date: 開始日期 (民國年格式，例如: 114/06/16)
            end_date: 結束日期 (民國年格式，例如: 114/07/16)

        Returns:
            除息資料DataFrame，失敗時返回None
        """
        # 轉換日期格式：民國年 114/06/16 -> 西元年 20250616
        def roc_to_ad_date(roc_date):
            try:
                year, month, day = roc_date.split('/')
                ad_year = int(year) + 1911
                return f"{ad_year}{month.zfill(2)}{day.zfill(2)}"
            except:
                return roc_date

        start_date_ad = roc_to_ad_date(start_date)
        end_date_ad = roc_to_ad_date(end_date)

        # 證交所API參數
        params = {
            'response': 'json',
            'startDate': start_date_ad,
            'endDate': end_date_ad
        }

        print(f"🔍 從證交所獲取除息資料...")
        print(f"📅 查詢期間: {start_date} ~ {end_date} (民國年)")
        print(f"📅 轉換期間: {start_date_ad} ~ {end_date_ad} (西元年)")
        print(f"🌐 URL: {self.tse_url}")

        try:
            # 發送請求
            response = requests.get(self.tse_url, params=params, timeout=30, verify=False)

            if response.status_code != 200:
                print(f"❌ 證交所HTTP請求失敗: {response.status_code}")
                return None

            # 解析JSON
            json_data = response.json()

            # 證交所的JSON結構可能不同，需要檢查
            if 'data' in json_data and json_data['data']:
                data = json_data['data']

                # 如果有欄位定義
                if 'fields' in json_data:
                    fields = json_data['fields']
                else:
                    # 使用預設欄位名稱
                    fields = ['除權息日期', '股票代號', '股票名稱', '除權息前收盤價',
                             '除權息參考價', '權值', '息值', '權值+息值', '權/息']

                # 創建DataFrame
                df = pd.DataFrame(data, columns=fields[:len(data[0])] if data else fields)
                df['資料來源'] = '證交所'
                print(f"✅ 證交所成功獲取 {len(df)} 筆除息資料")

                # 儲存證交所資料為Excel
                self._save_raw_data_to_excel(df, '證交所', start_date, end_date)

                return df
            else:
                print("❌ 證交所回應中無除息資料")
                return None

        except Exception as e:
            print(f"❌ 證交所獲取除息資料失敗: {e}")
            print(f"   可能原因: API格式變更或網路問題")
            return None

    def fetch_dividend_data(self, start_date: str, end_date: str) -> Optional[pd.DataFrame]:
        """
        從櫃買中心和證交所獲取除息資料

        Args:
            start_date: 開始日期 (民國年格式)
            end_date: 結束日期 (民國年格式)

        Returns:
            合併的除息資料DataFrame，失敗時返回None
        """
        print(f"🏦 開始獲取除息資料...")
        print(f"📅 查詢期間: {start_date} ~ {end_date}")
        print("=" * 60)

        all_dataframes = []

        # 1. 先從櫃買中心獲取資料
        tpex_df = self.fetch_dividend_data_from_tpex(start_date, end_date)
        if tpex_df is not None and not tpex_df.empty:
            all_dataframes.append(tpex_df)

        print()  # 空行分隔

        # 2. 再從證交所獲取資料
        tse_df = self.fetch_dividend_data_from_tse(start_date, end_date)
        if tse_df is not None and not tse_df.empty:
            all_dataframes.append(tse_df)

        # 3. 合併資料
        if not all_dataframes:
            print("❌ 無法從任何來源獲取除息資料")
            return None

        # 合併所有DataFrame
        combined_df = pd.concat(all_dataframes, ignore_index=True)

        print(f"\n📊 資料合併結果:")
        print(f"  櫃買中心: {len(tpex_df) if tpex_df is not None else 0} 筆")
        print(f"  證交所: {len(tse_df) if tse_df is not None else 0} 筆")
        print(f"  總計: {len(combined_df)} 筆")
        print(f"📁 原始資料已分別儲存為Excel檔案")

        return combined_df

    def _save_raw_data_to_excel(self, df: pd.DataFrame, source: str, start_date: str, end_date: str):
        """
        儲存原始除息資料為Excel檔案

        Args:
            df: 除息資料DataFrame
            source: 資料來源 ('櫃買中心' 或 '證交所')
            start_date: 開始日期
            end_date: 結束日期
        """
        if df is None or df.empty:
            print(f"⚠️  {source}無資料可儲存")
            return

        # 生成檔案名稱
        timestamp = datetime.now().strftime("%Y%m%d_%H%M")
        start_clean = start_date.replace('/', '')
        end_clean = end_date.replace('/', '')
        excel_file = f'{source}_除息資料_{start_clean}_{end_clean}_{timestamp}.xlsx'

        try:
            with pd.ExcelWriter(excel_file, engine='openpyxl') as writer:
                # 主要資料工作表
                df_main = df.copy()
                df_main.to_excel(writer, sheet_name='除息資料', index=False)

                # 摘要資訊工作表
                summary_data = {
                    '項目': [
                        '資料來源',
                        '查詢期間',
                        '資料筆數',
                        '欄位數量',
                        '產生時間',
                        'API網址'
                    ],
                    '內容': [
                        source,
                        f"{start_date} ~ {end_date}",
                        len(df),
                        len(df.columns),
                        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                        self.tpex_url if source == '櫃買中心' else self.tse_url
                    ]
                }
                summary_df = pd.DataFrame(summary_data)
                summary_df.to_excel(writer, sheet_name='摘要資訊', index=False)

                # 欄位說明工作表
                fields_data = {
                    '序號': range(1, len(df.columns) + 1),
                    '欄位名稱': df.columns.tolist(),
                    '資料類型': [str(df[col].dtype) for col in df.columns],
                    '非空值數量': [df[col].count() for col in df.columns],
                    '範例值': [str(df[col].iloc[0]) if len(df) > 0 else '' for col in df.columns]
                }
                fields_df = pd.DataFrame(fields_data)
                fields_df.to_excel(writer, sheet_name='欄位說明', index=False)

                # 統計分析工作表
                # 尋找配息相關欄位
                dividend_field = None
                for field in ['息值', '權值+息值', '現金股利', '配息金額']:
                    if field in df.columns:
                        dividend_field = field
                        break

                if dividend_field:
                    print(f"    📊 使用 '{dividend_field}' 欄位進行統計")
                    # 配息股票統計
                    dividend_stocks = df[df[dividend_field].astype(str) != '0.000000']
                    dividend_stocks = dividend_stocks[dividend_stocks[dividend_field].astype(str) != '0']

                    if not dividend_stocks.empty:
                        stats_data = []
                        for _, row in dividend_stocks.iterrows():
                            try:
                                # 使用找到的配息欄位
                                dividend_value = float(row[dividend_field])
                                if dividend_value > 0:
                                    # 如果使用的是權值+息值，嘗試取得純息值
                                    actual_dividend = dividend_value
                                    if dividend_field == '權值+息值' and '息值' in row.index:
                                        try:
                                            pure_dividend = float(row['息值'])
                                            if pure_dividend > 0:
                                                actual_dividend = pure_dividend
                                        except (ValueError, TypeError):
                                            pass  # 使用原始的權值+息值

                                    stats_data.append({
                                        '股票代碼': row.get('代號', row.get('股票代號', 'N/A')),
                                        '股票名稱': row.get('名稱', row.get('股票名稱', 'N/A')),
                                        '除息日期': row.get('除權息日期', 'N/A'),
                                        '每股配息': actual_dividend,
                                        '除息前價格': row.get('除權息前收盤價', 'N/A'),
                                        '配息欄位': dividend_field
                                    })
                            except (ValueError, TypeError):
                                continue

                        if stats_data:
                            stats_df = pd.DataFrame(stats_data)
                            stats_df = stats_df.sort_values('每股配息', ascending=False)
                            stats_df.to_excel(writer, sheet_name='配息統計', index=False)

            print(f"📁 {source}原始資料已儲存: {excel_file}")

            # 顯示統計資訊
            if dividend_field:
                dividend_count = 0
                try:
                    non_zero_dividends = df[df[dividend_field].astype(str) != '0.000000']
                    non_zero_dividends = non_zero_dividends[non_zero_dividends[dividend_field].astype(str) != '0']
                    dividend_count = len(non_zero_dividends)
                except:
                    pass
                print(f"📊 {source}統計: 總計 {len(df)} 筆，有配息 {dividend_count} 檔 (使用欄位: {dividend_field})")

        except Exception as e:
            print(f"❌ 儲存{source}Excel失敗: {e}")

    def get_inventory_stocks(self) -> Dict[str, Dict]:
        """
        從庫存Excel獲取所有持有的股票
        
        Returns:
            股票字典 {股票代碼: {股票名稱, 持有股數, ...}}
        """
        try:
            df_inventory = self.stock_system._read_inventory()
            
            if df_inventory.empty:
                print("📦 庫存為空")
                return {}
            
            stocks = {}
            for _, row in df_inventory.iterrows():
                stock_code = row['股票代碼']
                stocks[stock_code] = {
                    '股票名稱': row['股票名稱'],
                    '持有股數': int(row['持有股數']),
                    '平均成本': float(row['平均成本']),
                    '總成本': float(row['總成本'])
                }
            
            print(f"📊 從庫存中讀取到 {len(stocks)} 檔股票")
            return stocks
            
        except Exception as e:
            print(f"❌ 讀取庫存失敗: {e}")
            return {}
    
    def find_dividend_matches(self, dividend_df: pd.DataFrame, inventory_stocks: Dict[str, Dict]) -> List[Dict]:
        """
        找出庫存股票中有除息的項目

        Args:
            dividend_df: 除息資料DataFrame (包含櫃買中心和證交所資料)
            inventory_stocks: 庫存股票字典

        Returns:
            配息匹配列表
        """
        matches = []

        print(f"\n🔍 檢查庫存股票是否有除息...")
        print(f"📊 除息資料來源統計:")

        # 統計資料來源
        if '資料來源' in dividend_df.columns:
            source_counts = dividend_df['資料來源'].value_counts()
            for source, count in source_counts.items():
                print(f"  {source}: {count} 筆")

        print()

        for stock_code, stock_info in inventory_stocks.items():
            print(f"🔍 檢查 {stock_code} ({stock_info['股票名稱']})...")

            # 除錯：顯示除息資料的基本資訊
            print(f"  📊 除息資料總筆數: {len(dividend_df)}")
            print(f"  📋 除息資料欄位: {list(dividend_df.columns)}")

            # 在除息資料中搜尋此股票代碼
            # 支援不同的欄位名稱：代號、股票代號、股票代碼
            stock_rows = pd.DataFrame()

            # 嘗試不同的欄位名稱
            possible_code_fields = ['代號', '股票代號', '股票代碼', 'code', 'symbol']

            found_in_field = None
            for field in possible_code_fields:
                if field in dividend_df.columns:
                    print(f"  🔍 檢查欄位 '{field}'...")

                    # 除錯：顯示該欄位的一些範例值
                    sample_values = dividend_df[field].head(5).tolist()
                    print(f"    範例值: {sample_values}")

                    # 精確匹配
                    exact_match = dividend_df[dividend_df[field].astype(str) == stock_code]
                    if not exact_match.empty:
                        stock_rows = exact_match
                        found_in_field = field
                        print(f"  ✅ 在欄位 '{field}' 找到 {stock_code} (精確匹配)")
                        break

                    # 模糊匹配 (處理可能的空格或格式問題)
                    fuzzy_match = dividend_df[dividend_df[field].astype(str).str.strip() == stock_code]
                    if not fuzzy_match.empty:
                        stock_rows = fuzzy_match
                        found_in_field = field
                        print(f"  ✅ 在欄位 '{field}' 找到 {stock_code} (模糊匹配)")
                        break

                    print(f"    ➖ 欄位 '{field}' 無 {stock_code}")
                else:
                    print(f"  ⚠️  欄位 '{field}' 不存在")

            # 如果還是找不到，使用全欄位搜尋
            if stock_rows.empty:
                print(f"  🔍 在所有欄位中搜尋 {stock_code}...")
                stock_rows = dividend_df[dividend_df.apply(
                    lambda row: stock_code in str(row).replace(' ', ''), axis=1
                )]
                if not stock_rows.empty:
                    found_in_field = "全欄位搜尋"
                    print(f"  ✅ 在全欄位搜尋中找到 {stock_code}")
                else:
                    print(f"  ❌ 全欄位搜尋也未找到 {stock_code}")

                    # 除錯：檢查是否有類似的股票代碼
                    if len(stock_code) >= 4:
                        prefix = stock_code[:2]
                        similar_stocks = dividend_df[dividend_df.apply(
                            lambda row: prefix in str(row).replace(' ', ''), axis=1
                        )]
                        if not similar_stocks.empty:
                            print(f"  🔍 找到 {prefix} 開頭的股票: {len(similar_stocks)} 檔")
                            # 顯示前幾個作為參考
                            for field in possible_code_fields:
                                if field in similar_stocks.columns:
                                    sample_codes = similar_stocks[field].head(3).tolist()
                                    print(f"    範例 ({field}): {sample_codes}")
                                    break

            if not stock_rows.empty:
                print(f"✅ 找到 {stock_code} 的除息資料! (共 {len(stock_rows)} 筆)")

                for _, dividend_row in stock_rows.iterrows():
                    # 尋找配息金額欄位
                    dividend_amount = 0
                    dividend_field = ""
                    data_source = dividend_row.get('資料來源', '未知')

                    # 可能的配息欄位名稱 (按優先順序排列)
                    possible_fields = [
                        '息值',           # 櫃買中心主要使用
                        '權值+息值',      # 證交所可能使用 (包含權值和息值)
                        '現金股利',       # 通用欄位
                        '配息金額',       # 通用欄位
                        '股利',          # 簡化欄位
                        '配息',          # 簡化欄位
                        '現金配息'       # 完整描述
                    ]

                    print(f"    🔍 搜尋配息欄位，可用欄位: {list(dividend_row.index)}")

                    for field in possible_fields:
                        if field in dividend_row.index:
                            field_value = dividend_row[field]
                            print(f"    📊 檢查欄位 '{field}': '{field_value}' (類型: {type(field_value)})")

                            # 檢查欄位是否有值 (不是None, NaN, 空字串)
                            if pd.isna(field_value) or field_value == '' or field_value is None:
                                print(f"    ➖ 欄位 '{field}' 為空值，跳過")
                                continue

                            try:
                                # 特殊處理 '權值+息值' 欄位
                                if field == '權值+息值':
                                    print(f"    🎯 處理權值+息值欄位")

                                    # 先檢查是否有單獨的息值欄位且有值
                                    if '息值' in dividend_row.index:
                                        xi_value = dividend_row['息值']
                                        print(f"    📊 息值欄位內容: '{xi_value}' (類型: {type(xi_value)})")

                                        if not pd.isna(xi_value) and xi_value != '' and xi_value is not None:
                                            try:
                                                xi_amount = float(xi_value)
                                                if xi_amount > 0:
                                                    dividend_amount = xi_amount
                                                    dividend_field = '息值'
                                                    print(f"    ✅ 使用息值欄位: {dividend_amount}")
                                                else:
                                                    print(f"    ⚠️  息值為0，嘗試使用權值+息值")
                                                    dividend_amount = float(field_value)
                                                    dividend_field = field
                                                    print(f"    ✅ 使用權值+息值欄位: {dividend_amount}")
                                            except (ValueError, TypeError) as e:
                                                print(f"    ❌ 息值解析失敗: {e}，使用權值+息值")
                                                dividend_amount = float(field_value)
                                                dividend_field = field
                                                print(f"    ✅ 使用權值+息值欄位: {dividend_amount}")
                                        else:
                                            print(f"    ⚠️  息值欄位為空，使用權值+息值")
                                            dividend_amount = float(field_value)
                                            dividend_field = field
                                            print(f"    ✅ 使用權值+息值欄位: {dividend_amount}")
                                    else:
                                        print(f"    ⚠️  無息值欄位，直接使用權值+息值")
                                        dividend_amount = float(field_value)
                                        dividend_field = field
                                        print(f"    ✅ 使用權值+息值欄位: {dividend_amount}")
                                else:
                                    # 一般欄位處理
                                    dividend_amount = float(field_value)
                                    dividend_field = field
                                    print(f"    ✅ 使用{field}欄位: {dividend_amount}")

                                if dividend_amount > 0:
                                    print(f"    🎉 成功解析配息金額: {dividend_amount} (來源: {dividend_field})")
                                    break
                                else:
                                    print(f"    ⚠️  配息金額為0，繼續搜尋其他欄位")

                            except (ValueError, TypeError) as e:
                                print(f"    ❌ 欄位 '{field}' 解析失敗: {e}")
                                continue

                    if dividend_amount > 0:
                        # 計算配息收益
                        holdings = stock_info['持有股數']
                        total_dividend = holdings * dividend_amount

                        # 取得除息日期
                        dividend_date = dividend_row.get('除權息日期', '未知')

                        match_info = {
                            '股票代碼': stock_code,
                            '股票名稱': stock_info['股票名稱'],
                            '持有股數': holdings,
                            '配息金額': dividend_amount,
                            '配息欄位': dividend_field,
                            '總配息收益': total_dividend,
                            '除息日期': dividend_date,
                            '資料來源': data_source,
                            '除息資料': dividend_row.to_dict()
                        }

                        matches.append(match_info)

                        print(f"💰 {stock_code} 配息: {dividend_amount} 元/股 (來源: {data_source})")
                        print(f"💰 除息日期: {dividend_date}")
                        print(f"💰 持股 {holdings} 股，總收益: {total_dividend:.2f} 元")
                    else:
                        print(f"⚠️  {stock_code} 找到除息資料但無法解析配息金額 (來源: {data_source})")
            else:
                print(f"➖ {stock_code} 在櫃買中心和證交所都無除息資料")

        return matches
    
    def _initialize_dividend_database(self):
        """初始化配息資料庫Excel檔案"""
        dividend_db_file = "dividend_database.xlsx"

        if not os.path.exists(dividend_db_file):
            # 創建新的配息資料庫
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment

            wb = Workbook()
            ws = wb.active
            ws.title = "配息記錄"

            # 設定標題
            headers = [
                "年月", "股票代碼", "股票名稱", "持有股數",
                "每股配息", "總配息收益", "除息日期", "檢查日期", "資料來源", "備註"
            ]

            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")

            wb.save(dividend_db_file)
            print(f"✅ 已創建配息資料庫: {dividend_db_file}")

        return dividend_db_file

    def _read_dividend_database(self, db_file: str) -> pd.DataFrame:
        """讀取配息資料庫"""
        try:
            df = pd.read_excel(db_file, sheet_name="配息記錄")
            return df
        except Exception as e:
            print(f"❌ 讀取配息資料庫失敗: {e}")
            return pd.DataFrame()

    def _save_dividend_database(self, df: pd.DataFrame, db_file: str):
        """儲存配息資料庫"""
        try:
            with pd.ExcelWriter(db_file, mode='w') as writer:
                df.to_excel(writer, sheet_name="配息記錄", index=False)
            print("✅ 配息資料庫已更新")
        except Exception as e:
            print(f"❌ 儲存配息資料庫失敗: {e}")

    def _get_year_month_from_date_range(self, start_date: str, end_date: str) -> str:
        """
        從日期範圍取得年月字串

        Args:
            start_date: 開始日期 (民國年格式)
            end_date: 結束日期 (民國年格式)

        Returns:
            年月字串，例如: "2025-06"
        """
        try:
            # 解析民國年日期 (例如: "114/06/01")
            year_roc, month, _ = start_date.split('/')
            year_ad = int(year_roc) + 1911  # 轉換為西元年
            return f"{year_ad}-{month.zfill(2)}"
        except:
            return datetime.now().strftime("%Y-%m")

    def save_dividend_report(self, matches: List[Dict], start_date: str, end_date: str):
        """
        儲存配息報告到Excel資料庫

        Args:
            matches: 配息匹配列表
            start_date: 查詢開始日期
            end_date: 查詢結束日期
        """
        if not matches:
            print("📝 無配息資料需要儲存")
            return

        # 初始化配息資料庫
        db_file = self._initialize_dividend_database()

        # 讀取現有資料庫
        df_existing = self._read_dividend_database(db_file)

        # 取得年月
        year_month = self._get_year_month_from_date_range(start_date, end_date)
        print(f"📅 處理年月: {year_month}")

        # 準備新增的資料
        new_records = []
        updated_count = 0
        skipped_count = 0

        for match in matches:
            stock_code = match['股票代碼']

            # 檢查是否已存在相同年月和股票代碼的記錄
            existing_record = df_existing[
                (df_existing['年月'] == year_month) &
                (df_existing['股票代碼'] == stock_code)
            ]

            if not existing_record.empty:
                print(f"⚠️  {stock_code} 在 {year_month} 的配息記錄已存在，跳過新增")
                skipped_count += 1
                continue

            # 準備新記錄
            new_record = {
                '年月': year_month,
                '股票代碼': stock_code,
                '股票名稱': match['股票名稱'],
                '持有股數': match['持有股數'],
                '每股配息': match['配息金額'],
                '總配息收益': match['總配息收益'],
                '除息日期': match.get('除息日期', f"{start_date} ~ {end_date}"),
                '檢查日期': datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                '資料來源': match.get('資料來源', '未知'),
                '備註': f"來源欄位: {match['配息欄位']}"
            }

            new_records.append(new_record)
            updated_count += 1
            print(f"✅ 準備新增 {stock_code} 的配息記錄")

        # 如果有新記錄需要新增
        if new_records:
            df_new = pd.DataFrame(new_records)

            if df_existing.empty:
                df_updated = df_new
            else:
                df_updated = pd.concat([df_existing, df_new], ignore_index=True)

            # 按年月和股票代碼排序
            df_updated = df_updated.sort_values(['年月', '股票代碼'])

            # 儲存更新後的資料庫
            self._save_dividend_database(df_updated, db_file)

            print(f"📊 配息資料庫更新完成:")
            print(f"  ✅ 新增記錄: {updated_count} 筆")
            print(f"  ⚠️  跳過記錄: {skipped_count} 筆")
            print(f"  📁 資料庫檔案: {db_file}")
        else:
            print(f"📊 無新記錄需要新增 (跳過 {skipped_count} 筆重複記錄)")

        # 計算總收益
        total_income = sum(match['總配息收益'] for match in matches)
        print(f"💰 本期總配息收益: {total_income:.2f} 元")

        # 同時儲存CSV報告 (保留原有功能)
        self._save_csv_report(matches, start_date, end_date)

    def _save_csv_report(self, matches: List[Dict], start_date: str, end_date: str):
        """儲存CSV格式的配息報告 (保留原有功能)"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M")

        # 創建配息報告DataFrame
        report_data = []
        for match in matches:
            report_data.append({
                '檢查日期': datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                '查詢期間': f"{start_date} ~ {end_date}",
                '股票代碼': match['股票代碼'],
                '股票名稱': match['股票名稱'],
                '持有股數': match['持有股數'],
                '每股配息': match['配息金額'],
                '總配息收益': match['總配息收益'],
                '除息日期': match.get('除息日期', '未知'),
                '資料來源': match.get('資料來源', '未知'),
                '配息欄位': match['配息欄位']
            })

        df_report = pd.DataFrame(report_data)

        # 儲存CSV報告
        report_file = f"配息報告_{timestamp}.csv"
        df_report.to_csv(report_file, index=False, encoding='utf-8-sig')

        print(f"📄 CSV報告已儲存: {report_file}")

    def show_dividend_database(self, year_month: str = None, stock_code: str = None):
        """
        顯示配息資料庫內容

        Args:
            year_month: 指定年月，例如 "2025-06"，None表示顯示全部
            stock_code: 指定股票代碼，None表示顯示全部
        """
        db_file = "dividend_database.xlsx"

        if not os.path.exists(db_file):
            print("📦 配息資料庫不存在")
            return

        df = self._read_dividend_database(db_file)

        if df.empty:
            print("📦 配息資料庫為空")
            return

        # 篩選條件
        filtered_df = df.copy()

        if year_month:
            filtered_df = filtered_df[filtered_df['年月'] == year_month]
            if filtered_df.empty:
                print(f"📦 找不到 {year_month} 的配息記錄")
                return

        if stock_code:
            filtered_df = filtered_df[filtered_df['股票代碼'] == stock_code]
            if filtered_df.empty:
                print(f"📦 找不到 {stock_code} 的配息記錄")
                return

        print(f"\n💰 配息資料庫 (共 {len(filtered_df)} 筆記錄)")
        print("=" * 80)

        # 按年月分組顯示
        for year_month_group in filtered_df['年月'].unique():
            month_data = filtered_df[filtered_df['年月'] == year_month_group]
            month_total = month_data['總配息收益'].sum()

            print(f"\n📅 {year_month_group} (共 {len(month_data)} 檔股票，總收益: {month_total:.2f} 元)")
            print("-" * 60)

            for _, row in month_data.iterrows():
                print(f"  {row['股票代碼']} ({row['股票名稱']})")
                print(f"    持股: {row['持有股數']:,} 股")
                print(f"    配息: {row['每股配息']} 元/股")
                print(f"    收益: {row['總配息收益']:.2f} 元")
                print(f"    日期: {row['除息日期']}")
                if '資料來源' in row.index:
                    print(f"    來源: {row['資料來源']}")
                if row['備註']:
                    print(f"    備註: {row['備註']}")
                print()

        # 顯示統計資訊
        total_income = filtered_df['總配息收益'].sum()
        unique_stocks = filtered_df['股票代碼'].nunique()
        unique_months = filtered_df['年月'].nunique()

        print(f"📊 統計資訊:")
        print(f"  💰 總配息收益: {total_income:.2f} 元")
        print(f"  📈 涉及股票: {unique_stocks} 檔")
        print(f"  📅 涉及月份: {unique_months} 個月")

    def get_dividend_summary(self, year: str = None) -> pd.DataFrame:
        """
        取得配息摘要統計

        Args:
            year: 指定年份，例如 "2025"，None表示全部年份

        Returns:
            配息摘要DataFrame
        """
        db_file = "dividend_database.xlsx"

        if not os.path.exists(db_file):
            print("📦 配息資料庫不存在")
            return pd.DataFrame()

        df = self._read_dividend_database(db_file)

        if df.empty:
            print("📦 配息資料庫為空")
            return pd.DataFrame()

        # 篩選年份
        if year:
            df = df[df['年月'].str.startswith(year)]
            if df.empty:
                print(f"📦 找不到 {year} 年的配息記錄")
                return pd.DataFrame()

        # 按年月和股票代碼分組統計
        summary = df.groupby(['年月', '股票代碼', '股票名稱']).agg({
            '持有股數': 'first',
            '每股配息': 'sum',  # 如果同一個月有多次配息
            '總配息收益': 'sum'
        }).reset_index()

        return summary

    def run_daily_check(self):
        """執行每日除息檢查"""
        print("🏦 每日除息資料檢查器 (雙重來源)")
        print("=" * 60)
        print(f"🕐 執行時間: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        print(f"📊 資料來源: 櫃買中心 + 證交所")

        # 1. 取得查詢日期範圍
        start_date, end_date = self.get_date_range_roc()
        print(f"📅 查詢前一個月: {start_date} ~ {end_date}")

        # 2. 獲取庫存股票
        inventory_stocks = self.get_inventory_stocks()
        if not inventory_stocks:
            print("❌ 無庫存股票，結束檢查")
            return

        # 3. 獲取除息資料 (櫃買中心 + 證交所)
        dividend_df = self.fetch_dividend_data(start_date, end_date)
        if dividend_df is None or dividend_df.empty:
            print("❌ 無法從任何來源獲取除息資料，結束檢查")
            return

        # 4. 找出配息匹配
        matches = self.find_dividend_matches(dividend_df, inventory_stocks)

        # 5. 顯示結果
        print(f"\n📊 檢查結果:")
        print("=" * 40)

        if matches:
            print(f"✅ 找到 {len(matches)} 檔股票有配息!")

            # 按資料來源分組顯示
            tpex_matches = [m for m in matches if m.get('資料來源') == '櫃買中心']
            tse_matches = [m for m in matches if m.get('資料來源') == '證交所']

            if tpex_matches:
                print(f"\n📈 櫃買中心 ({len(tpex_matches)} 檔):")
                for match in tpex_matches:
                    print(f"  💰 {match['股票代碼']} ({match['股票名稱']}): "
                          f"{match['配息金額']} 元/股 × {match['持有股數']} 股 = "
                          f"{match['總配息收益']:.2f} 元")

            if tse_matches:
                print(f"\n🏛️ 證交所 ({len(tse_matches)} 檔):")
                for match in tse_matches:
                    print(f"  💰 {match['股票代碼']} ({match['股票名稱']}): "
                          f"{match['配息金額']} 元/股 × {match['持有股數']} 股 = "
                          f"{match['總配息收益']:.2f} 元")

            # 計算總收益
            total_income = sum(match['總配息收益'] for match in matches)
            print(f"\n💰 總配息收益: {total_income:.2f} 元")

            # 6. 儲存報告
            self.save_dividend_report(matches, start_date, end_date)
        else:
            print("➖ 本期無庫存股票配息")

        print(f"\n✅ 每日檢查完成!")


def main():
    """主程序 - 命令列介面"""
    checker = DailyDividendChecker()

    print("🏦 每日除息檢查器")
    print("=" * 50)

    while True:
        print("\n請選擇操作:")
        print("1. 執行每日除息檢查")
        print("2. 查看配息資料庫")
        print("3. 查看特定年月配息")
        print("4. 查看特定股票配息")
        print("5. 配息年度摘要")
        print("6. 退出")

        choice = input("\n請輸入選項 (1-6): ").strip()

        if choice == "1":
            # 執行每日檢查
            checker.run_daily_check()

        elif choice == "2":
            # 查看完整配息資料庫
            checker.show_dividend_database()

        elif choice == "3":
            # 查看特定年月配息
            year_month = input("請輸入年月 (格式: 2025-06): ").strip()
            if year_month:
                checker.show_dividend_database(year_month=year_month)
            else:
                print("❌ 請輸入有效的年月格式")

        elif choice == "4":
            # 查看特定股票配息
            stock_code = input("請輸入股票代碼: ").strip().upper()
            if stock_code:
                checker.show_dividend_database(stock_code=stock_code)
            else:
                print("❌ 請輸入股票代碼")

        elif choice == "5":
            # 配息年度摘要
            year = input("請輸入年份 (格式: 2025，留空顯示全部): ").strip()
            summary_df = checker.get_dividend_summary(year if year else None)

            if not summary_df.empty:
                print(f"\n📊 配息摘要 {'(' + year + '年)' if year else '(全部年份)'}")
                print("=" * 80)

                # 按年月分組顯示
                for year_month in summary_df['年月'].unique():
                    month_data = summary_df[summary_df['年月'] == year_month]
                    month_total = month_data['總配息收益'].sum()

                    print(f"\n📅 {year_month} - 總收益: {month_total:.2f} 元")
                    print("-" * 40)

                    for _, row in month_data.iterrows():
                        print(f"  {row['股票代碼']} ({row['股票名稱']}): {row['總配息收益']:.2f} 元")

                # 總計
                total = summary_df['總配息收益'].sum()
                print(f"\n💰 總配息收益: {total:.2f} 元")

        elif choice == "6":
            print("👋 感謝使用每日除息檢查器！")
            break

        else:
            print("❌ 無效的選項，請重新選擇")


def run_auto_check():
    """自動執行每日檢查 (用於批次檔案或排程器)"""
    checker = DailyDividendChecker()
    checker.run_daily_check()


if __name__ == "__main__":
    import sys

    # 如果有命令列參數 --auto，則直接執行檢查
    if len(sys.argv) > 1 and sys.argv[1] == "--auto":
        run_auto_check()
    else:
        main()
